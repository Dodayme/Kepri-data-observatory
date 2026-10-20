"""Audit read-only berkas referensi; tidak menyunting workbook."""
from pathlib import Path
import json
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(__file__).resolve().parents[1]
book = next(root.glob('Data_Sains*.xlsx'))
wb = load_workbook(book, read_only=True, data_only=True)
sheet = wb['Master_Integrasi']
values = list(sheet.values)
rows = [dict(zip(values[0], row)) for row in values[1:] if any(v is not None for v in row)]
audit = []
for row in rows:
    diff = row['Kes_Puskesmas'] - row['Kes_PuskesmasRawatInap'] - row['Kes_PuskesmasNonRawatInap']
    if diff:
        audit.append({k: row[k] for k in ['Kabupaten_Kota', 'Tahun', 'Kes_Puskesmas', 'Kes_PuskesmasRawatInap', 'Kes_PuskesmasNonRawatInap']} | {'selisih': diff})
report = {'file': book.name, 'sheets': wb.sheetnames, 'columns': list(values[0]), 'rows': len(rows), 'regions': sorted({r['Kabupaten_Kota'] for r in rows}), 'years': sorted({r['Tahun'] for r in rows}), 'puskesmas_mismatch_count': len(audit), 'puskesmas_mismatches': audit}
(root/'docs/reference-inspection.json').write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding='utf-8')
pdf = next(root.glob('Laporan*.pdf'))
(root/'docs/report-extracted.txt').write_text('\n'.join(p.extract_text() or '' for p in PdfReader(pdf).pages), encoding='utf-8')
print(json.dumps({k: v for k, v in report.items() if k != 'puskesmas_mismatches'}, indent=2))
